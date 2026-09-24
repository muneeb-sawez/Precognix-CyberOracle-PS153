import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    """Sets background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a cell."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    """Sets subtle borders for a table."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="none"/>\n'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideV w:val="none"/>\n'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def build_document():
    doc = Document()
    
    # Set Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(51, 65, 85) # Slate 700

    # Color Palette
    PRIMARY_NAVY = RGBColor(15, 23, 42)    # #0F172A
    CYAN_HEADER = RGBColor(14, 116, 144)   # #0E7490
    DARK_SLATE = RGBColor(30, 41, 59)     # #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139)   # #64748B
    ACCENT_BLUE = RGBColor(2, 132, 199)    # #0284C7
    
    # -------------------------------------------------------------
    # DOCUMENT TITLE & HEADER BLOCK
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("CYBER ORACLE // SIH 2026 INTERNAL ROUND")
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_NAVY

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(16)
    run_sub = subtitle_p.add_run("Proactive Cyber Threat Forecasting & Mitigation Prototype — Technical Architecture & Viva Guide")
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = CYAN_HEADER

    # Meta Info Card / Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    meta_data = [
        [("Problem Statement ID:", " 26153"), ("Project Name:", " Cyber Oracle")],
        [("Target Domain:", " Proactive Threat Forecasting & Mitigation"), ("Tech Stack:", " React 19, Vite, Tailwind CSS v4, Recharts, Framer Motion")]
    ]
    
    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            set_cell_background(cell, "F1F5F9") # Slate 100
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            cell.width = Inches(3.4)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            label, val = meta_data[row_idx][col_idx]
            r1 = p.add_run(label)
            r1.font.bold = True
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = DARK_SLATE
            r2 = p.add_run(val)
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = PRIMARY_NAVY

    set_table_borders(meta_table, color="CBD5E1", sz="4")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper function for headings
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = CYAN_HEADER
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = DARK_SLATE
        return p

    def add_bullet(p_or_str, bold_prefix="", text=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.bold = True
            r_bold.font.color.rgb = DARK_SLATE
        r_text = p.add_run(text)
        r_text.font.color.rgb = RGBColor(51, 65, 85)
        return p

    # -------------------------------------------------------------
    # SECTION 1: TECHNOLOGIES USED
    # -------------------------------------------------------------
    add_h1("1. Technologies Used in the Prototype")

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(8)
    p_intro.add_run("The Cyber Oracle prototype is engineered as a modern, high-performance, real-time threat intelligence web interface. Below is the detailed breakdown of the complete technology stack powering the application:")

    tech_table = doc.add_table(rows=1, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_table.autofit = False

    # Header Row
    hdr_cells = tech_table.rows[0].cells
    headers = ["Category", "Technology / Library", "Role & Function in Prototype"]
    col_widths = [Inches(1.5), Inches(2.0), Inches(3.3)]
    
    for i, title in enumerate(headers):
        set_cell_background(hdr_cells[i], "0F172A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)

    tech_data = [
        ("Frontend Core", "React 19 (`react`, `react-dom`)", "Provides declarative component-based UI rendering, custom Hooks for simulation state management, and fast virtual DOM updates."),
        ("Build Tool & Server", "Vite 8.2 (`@vitejs/plugin-react`)", "Ultra-fast bundler and development server with instant Hot Module Replacement (HMR) and optimized ES module compilation."),
        ("Styling & Design System", "Tailwind CSS v4 (`@tailwindcss/vite`)", "Utility-first CSS styling framework enabling dark/light mode switching, glassmorphism card aesthetics, and responsive layout grids."),
        ("Data Visualization", "Recharts v3 (`recharts`)", "Renders interactive time-series risk forecasting graphs, historical risk trend lines, and confidence interval shaded bands."),
        ("Animations & Motion", "Framer Motion v13 (`framer-motion`)", "Delivers smooth tab switching transitions, animated risk status badges, and interactive telemetry updates."),
        ("Iconography", "Lucide React (`lucide-react`)", "Provides crisp, modern vector icons for threat indicators, navigation tabs, attack vectors, and terminal controls."),
        ("State & Style Utilities", "clsx & tailwind-merge", "Handles dynamic conditional CSS class merging for active attack alerts, risk tier badges, and theme changes."),
        ("Interaction Feedback", "canvas-confetti", "Provides instant visual confirmation upon executing automated threat mitigation protocols."),
        ("Code Quality & Linting", "Oxlint (`oxlint`)", "High-speed Rust-based linter enforcing strict code quality, clean syntax, and preventing runtime React state bugs."),
        ("Benchmark Data Source", "CIC-IDS2018 Telemetry Dataset", "Serves as the empirical statistical baseline for attack scenario parameters (packet rates, flag entropy, SHAP feature vectors).")
    ]

    for row_idx, (cat, tech, role) in enumerate(tech_data):
        row_cells = tech_table.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for i, text in enumerate([cat, tech, role]):
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top=80, bottom=80, left=120, right=120)
            row_cells[i].width = col_widths[i]
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if i == 1:
                r.font.bold = True
                r.font.color.rgb = ACCENT_BLUE

    set_table_borders(tech_table, color="E2E8F0", sz="4")
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 2: HOW THE PROTOTYPE IS BUILT
    # -------------------------------------------------------------
    add_h1("2. How the Prototype is Built")

    p_build = doc.add_paragraph()
    p_build.paragraph_format.space_after = Pt(8)
    p_build.add_run("The Cyber Oracle prototype is built as a single-page application (SPA) featuring a reactive state-driven simulation engine. The key architectural highlights are outlined below:")

    add_h2("A. Application Architecture & File Structure")
    add_bullet("", "src/App.jsx: ", "The central orchestration container managing application-wide state (active scenario, timeline step index, playback speed, theme, active tab, and mitigation state).")
    add_bullet("", "src/data/attackScenarios.js: ", "The core mathematical telemetry simulation engine. Generates realistic live network metrics, SHAP feature attributions, MITRE ATT&CK technique stages, and network topology node threat levels based on the CIC-IDS2018 benchmark.")
    add_bullet("", "src/components/Header.jsx: ", "Top navigation bar containing scenario selectors, live play/pause controls, mitigation toggles, theme switcher (Dark/Light mode), and tab navigation.")
    add_bullet("", "src/components/HeroForecastWidget.jsx: ", "Renders the flagship K-Step Forecast card featuring radial risk gauges, risk tier badges (Nominal, Elevated, Critical), and time-series line charts.")
    add_bullet("", "src/components/MitreMatrix.jsx: ", "Displays real-time stage progression across the MITRE ATT&CK framework (Reconnaissance → Exploitation → Lateral Movement → C2 → Exfiltration).")
    add_bullet("", "src/components/ExplainableShapCard.jsx: ", "Provides Explainable AI (XAI) feature importance rankings showing exactly which packet parameters (e.g., TCP flag entropy, destination port 445 traffic) triggered the alert.")
    add_bullet("", "src/components/NetworkTopologyMap.jsx & TelemetryTerminal.jsx: ", "Displays active network nodes/links and streams live packet log telemetry in console format.")
    add_bullet("", "src/components/views/*: ", "Modular dedicated views including Topology Explorer, Mitigation Center, Analytics History, and Ingestion Settings.")

    add_h2("B. Live Simulation & Time-Scrubber Loop")
    p_sim = doc.add_paragraph()
    p_sim.paragraph_format.space_after = Pt(6)
    p_sim.paragraph_format.line_spacing = 1.15
    p_sim.add_run("The prototype implements a reactive event loop using React's ")
    r_code = p_sim.add_run("useEffect")
    r_code.font.bold = True
    p_sim.add_run(" hook. When play mode is active, an internal timer ticks every ")
    p_sim.add_run("1.2 seconds / playbackSpeed").bold = True
    p_sim.add_run(", dynamically incrementing the ")
    p_sim.add_run("stepIndex").bold = True
    p_sim.add_run(". Each tick triggers ")
    p_sim.add_run("getStepData()").bold = True
    p_sim.add_run(" to compute updated risk curves, MITRE ATT&CK statuses, node colors, and packet logs seamlessly.")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # SECTION 3: FRONTEND AND BACKEND CONNECTION
    # -------------------------------------------------------------
    add_h1("3. How Frontend and Backend are Connected")

    p_conn_intro = doc.add_paragraph()
    p_conn_intro.paragraph_format.space_after = Pt(8)
    p_conn_intro.add_run("In a production cybersecurity platform, connection between frontend and backend requires ultra-low latency telemetry streaming and asynchronous REST endpoints. Here is how it operates in the prototype vs. the production architecture:")

    add_h2("A. Prototype Architecture (Client-Side Simulation Layer)")
    add_bullet("", "In-Memory API Proxy: ", "In this frontend prototype, `src/data/attackScenarios.js` acts as an in-memory mock API backend. `getStepData(scenarioId, stepIndex, isMitigated)` simulates REST API responses by computing exact JSON data payloads synchronously.")
    add_bullet("", "PCAP Upload Simulation: ", "The PCAP Upload Modal allows users to upload local `.pcap` files. The frontend parses file metadata, simulates backend packet ingestion, and automatically switches the active view to the threat dashboard.")

    add_h2("B. Production Connection Model (Full Architecture Specs)")
    p_prod = doc.add_paragraph()
    p_prod.paragraph_format.space_after = Pt(6)
    p_prod.add_run("For full enterprise deployment, the frontend connects to the backend services via two main communication protocols:")

    conn_table = doc.add_table(rows=1, cols=3)
    conn_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    conn_table.autofit = False

    c_hdr = conn_table.rows[0].cells
    c_headers = ["Protocol / Connection", "Backend Component", "Data & Usage Description"]
    c_widths = [Inches(1.8), Inches(2.0), Inches(3.0)]

    for i, title in enumerate(c_headers):
        set_cell_background(c_hdr[i], "0F172A")
        set_cell_margins(c_hdr[i], top=100, bottom=100, left=120, right=120)
        c_hdr[i].width = c_widths[i]
        p = c_hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    conn_data = [
        ("WebSockets (`wss://`)", "Python FastAPI / Kafka Stream Consumer", "Bi-directional, persistent stream delivering live network packet logs, instant threat alerts, and real-time risk scores to the frontend terminal without polling."),
        ("REST API (`POST /api/mitigate`)", "Mitigation Engine (FastAPI)", "Sends firewall rule execution payloads (e.g., block IP, isolate subnet via iptables/BGP) triggered when user toggles Automated Mitigation."),
        ("REST API (`POST /api/ingest/pcap`)", "Packet Parser (dpkt / Scapy Engine)", "Handles PCAP file upload, processes raw PCAP byte streams, extracts feature vectors, and feeds them into the ML inference model."),
        ("REST API (`GET /api/analytics`)", "Database (PostgreSQL / TimescaleDB)", "Fetches historical attack logs, past incident trends, and benchmark comparisons for the Analytics & History view.")
    ]

    for row_idx, (proto, comp, desc) in enumerate(conn_data):
        r_cells = conn_table.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for i, text in enumerate([proto, comp, desc]):
            set_cell_background(r_cells[i], bg_color)
            set_cell_margins(r_cells[i], top=80, bottom=80, left=120, right=120)
            r_cells[i].width = c_widths[i]
            p = r_cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if i == 0:
                r.font.bold = True
                r.font.color.rgb = ACCENT_BLUE

    set_table_borders(conn_table, color="E2E8F0", sz="4")
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # -------------------------------------------------------------
    # SECTION 4: MOST EXPECTED VIVA QUESTIONS & BEST ANSWERS
    # -------------------------------------------------------------
    add_h1("4. Most Expected Viva Questions & Best Answers")

    p_viva_intro = doc.add_paragraph()
    p_viva_intro.paragraph_format.space_after = Pt(8)
    p_viva_intro.add_run("Below are the exact top 10 questions expected from judges during the SIH internal presentation, paired with concise, pin-point, high-impact answers:")

    viva_qna = [
        (
            "Q1: What exact problem does Cyber Oracle solve?",
            "Cyber Oracle shifts cybersecurity from reactive breach detection to proactive threat forecasting. It predicts attack infiltration probability K-steps ahead into the future, enabling automated mitigation before critical asset compromise occurs."
        ),
        (
            "Q2: What baseline dataset is used for threat modeling in this project?",
            "We base our telemetry metrics on the CIC-IDS2018 benchmark dataset, simulating real-world network parameters such as flow duration, packet arrival rates, TCP flag entropy, and payload size."
        ),
        (
            "Q3: How does the K-Step Threat Forecasting engine work?",
            "Time-series ML models (such as LSTM/XGBoost) analyze rolling historical telemetry windows to project future attack escalation curves and output a percentage probability of imminent network breach."
        ),
        (
            "Q4: What is SHAP and why is Explainable AI (XAI) crucial here?",
            "SHAP (SHapley Additive exPlanations) identifies exact packet features (e.g., high TCP SYN flag entropy or destination port 445 traffic) driving the risk score, eliminating 'black-box' AI ambiguity so security analysts can verify alerts instantly."
        ),
        (
            "Q5: How are the Frontend and Backend connected in production?",
            "WebSockets (`wss://`) stream live telemetry logs and real-time risk scores from a Python FastAPI backend, while REST APIs handle asynchronous tasks like PCAP file uploads, historical analytics queries, and automated firewall rule execution."
        ),
        (
            "Q6: How does the Automated AI Mitigation feature operate?",
            "When triggered, the system simulates immediate active network isolation (reducing risk score by ~80%+), isolates compromised subnets, and generates actionable defense rules (such as `iptables` IP blocking or BGP blackholing)."
        ),
        (
            "Q7: Why did you choose React 19, Vite, and Tailwind CSS v4 for the frontend?",
            "React 19 ensures fast component-based UI state updates; Vite provides instant HMR and optimized bundling; Tailwind CSS v4 enables lightweight utility-first styling for dark-mode glassmorphism aesthetics and responsive performance."
        ),
        (
            "Q8: How does the prototype handle PCAP file uploads?",
            "The upload modal parses `.pcap` packet capture metadata, feeds packet headers into the local telemetry processing state, and dynamically updates dashboard widgets to demonstrate live file ingestion."
        ),
        (
            "Q9: What differentiates Cyber Oracle from standard SIEM tools like Splunk or Sentinel?",
            "Standard SIEMs generate alerts after logs indicate a breach has occurred. Cyber Oracle forecasts threats before exploitation (proactive), provides SHAP feature attributions (explainable), and triggers active defense protocols (automated)."
        ),
        (
            "Q10: How does Cyber Oracle scale for high-throughput enterprise network traffic?",
            "In production, raw packet streaming is ingested via Apache Kafka event brokers and processed by lightweight C++/Python parsing microservices, pushing pre-aggregated inference results over WebSockets to maintain zero UI latency."
        )
    ]

    for idx, (q, a) in enumerate(viva_qna, 1):
        # Q&A Box / Container
        qa_table = doc.add_table(rows=2, cols=1)
        qa_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        qa_table.autofit = False

        # Header cell (Question)
        cell_q = qa_table.rows[0].cells[0]
        set_cell_background(cell_q, "1E293B") # Dark Slate
        set_cell_margins(cell_q, top=70, bottom=70, left=100, right=100)
        cell_q.width = Inches(6.8)
        pq = cell_q.paragraphs[0]
        pq.paragraph_format.space_after = Pt(0)
        rq = pq.add_run(q)
        rq.font.bold = True
        rq.font.size = Pt(10)
        rq.font.color.rgb = RGBColor(255, 255, 255)

        # Body cell (Answer)
        cell_a = qa_table.rows[1].cells[0]
        set_cell_background(cell_a, "F8FAFC") # Light slate
        set_cell_margins(cell_a, top=80, bottom=80, left=100, right=100)
        cell_a.width = Inches(6.8)
        pa = cell_a.paragraphs[0]
        pa.paragraph_format.space_after = Pt(0)
        pa.paragraph_format.line_spacing = 1.15
        ra_label = pa.add_run("Pin-Point Answer: ")
        ra_label.font.bold = True
        ra_label.font.size = Pt(9.5)
        ra_label.font.color.rgb = CYAN_HEADER
        ra_text = pa.add_run(a)
        ra_text.font.size = Pt(9.5)
        ra_text.font.color.rgb = PRIMARY_NAVY

        set_table_borders(qa_table, color="CBD5E1", sz="4")
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Save Document
    import os
    filename = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Cyber_Oracle_SIH_Presentation_Guide.docx")
    doc.save(filename)
    print(f"Document successfully created and saved to: {filename}")

if __name__ == "__main__":
    build_document()
